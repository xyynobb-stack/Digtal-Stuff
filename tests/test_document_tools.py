import importlib.util
import json
import tempfile
import threading
import unittest
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path


SCRIPT = Path(__file__).parents[1] / "resources" / "feature-workers" / "document_tools.py"
SPEC = importlib.util.spec_from_file_location("document_tools", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(MODULE)


class DocumentToolsTests(unittest.TestCase):
    # @lat: [[feature-workspace#OCR#MinerU page routing]]
    def test_mineru_response_and_structured_tables_are_normalized(self):
        markdown = MODULE._extract_mineru_markdown(
            {
                "results": {
                    "page-1.png": {
                        "md_content": (
                            "第一条 服务\n\n"
                            "<table><tr><th colspan='2'>费用</th></tr>"
                            "<tr><td rowspan='2'>软件</td><td>100元</td></tr>"
                            "<tr><td>200元</td></tr></table>"
                        )
                    }
                }
            },
            "page-1.png",
        )
        blocks = MODULE._structured_text_blocks(
            markdown,
            "old",
            {"paragraph": 0, "table": 0},
        )

        self.assertEqual(blocks[0]["type"], "paragraph")
        table = blocks[1]
        self.assertEqual(table["type"], "table")
        self.assertEqual(table["rows"][0]["cells"][0]["colSpan"], 2)
        self.assertEqual(table["rows"][1]["cells"][0]["rowSpan"], 2)
        self.assertEqual(table["rows"][2]["cells"][0]["column"], 2)

    def test_image_ocr_uses_mineru_before_local_runtime(self):
        original_config = MODULE._mineru_config
        original_parse = MODULE._mineru_parse_image
        MODULE._mineru_config = lambda: {
            "enabled": True,
            "apiBase": "http://mineru.invalid",
        }
        MODULE._mineru_parse_image = lambda content, filename, config: (
            "## 委托事项\n\n"
            "![](images/server-only.jpg)\n\n"
            "<table><tr><td>序号</td><td>项目</td></tr>"
            "<tr><td>1</td><td>MinerU 识别结果</td></tr>"
            "<tr><td colspan='2'>合计</td></tr></table>"
        )
        try:
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "scan.png"
                path.write_bytes(b"test-image")
                result = MODULE._ocr({"path": str(path)})
        finally:
            MODULE._mineru_config = original_config
            MODULE._mineru_parse_image = original_parse

        self.assertEqual(result["pages"][0]["source"], "mineru")
        self.assertNotIn("images/", result["text"])
        self.assertNotIn("<table>", result["text"])
        self.assertIn("MinerU 识别结果", result["text"])
        blocks = result["pages"][0]["blocks"]
        self.assertEqual(blocks[0]["style"], "heading")
        self.assertEqual(blocks[0]["text"], "委托事项")
        self.assertEqual(blocks[1]["type"], "table")
        self.assertEqual(blocks[1]["rows"][2]["cells"][0]["colSpan"], 2)

    def test_mineru_client_posts_expected_multipart_request(self):
        captured = {}

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                length = int(self.headers.get("Content-Length", "0"))
                captured["path"] = self.path
                captured["content_type"] = self.headers.get("Content-Type")
                captured["body"] = self.rfile.read(length)
                payload = {
                    "results": {
                        "page.png": {"markdown": "服务器识别结果"}
                    }
                }
                encoded = json.dumps(payload).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(encoded)))
                self.end_headers()
                self.wfile.write(encoded)

            def log_message(self, format, *args):
                return

        server = HTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            result = MODULE._mineru_parse_image(
                b"image-bytes",
                "page.png",
                {
                    "apiBase": f"http://127.0.0.1:{server.server_port}",
                    "pageTimeoutSeconds": 2,
                },
            )
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=2)

        self.assertEqual(result, "服务器识别结果")
        self.assertEqual(captured["path"], "/file_parse")
        self.assertIn("multipart/form-data", captured["content_type"])
        self.assertIn(b'name="table_enable"', captured["body"])
        self.assertIn(b'name="files"', captured["body"])

    @unittest.skipUnless(
        importlib.util.find_spec("fitz"),
        "PyMuPDF is available in the packaged feature runtime",
    )
    def test_scanned_pdf_page_is_routed_to_mineru(self):
        import fitz

        original_config = MODULE._mineru_config
        original_parse = MODULE._mineru_parse_image
        MODULE._mineru_config = lambda: {
            "enabled": True,
            "apiBase": "http://mineru.invalid",
            "maxConcurrency": 1,
            "renderDpi": 144,
        }
        MODULE._mineru_parse_image = (
            lambda content, filename, config: "扫描页结构化结果"
        )
        try:
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "scan.pdf"
                document = fitz.open()
                page = document.new_page(width=200, height=200)
                image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
                image.clear_with(255)
                page.insert_image(page.rect, stream=image.tobytes("png"))
                document.save(path)
                document.close()
                pages = MODULE._pdf_page_results(path)
        finally:
            MODULE._mineru_config = original_config
            MODULE._mineru_parse_image = original_parse

        self.assertEqual(pages[0]["source"], "mineru")
        self.assertEqual(pages[0]["text"], "扫描页结构化结果")

    def test_visual_pdf_lines_are_merged_into_contract_clauses(self):
        lines = [
            "第一条 服务",
            "服务范围包括系统建设。",
            "第二条 费用",
            "合同费用为人民币 100 元。",
        ]

        self.assertEqual(
            MODULE._merge_text_lines(lines),
            [
                "第一条 服务 服务范围包括系统建设。",
                "第二条 费用 合同费用为人民币 100 元。",
            ],
        )

    # @lat: [[feature-workspace#Contract comparison#Deterministic paragraph diff]]
    def test_compare_classifies_added_removed_and_modified_paragraphs(self):
        documents = {
            "old.docx": ["第一条 服务范围", "第二条 金额为100元", "删除条款"],
            "new.docx": ["第一条 服务范围", "第二条 金额为200元", "新增条款"],
        }
        original = MODULE._text_paragraphs
        MODULE._text_paragraphs = lambda path: documents[path.name]
        try:
            result = MODULE._compare({"oldPath": "old.docx", "newPath": "new.docx"})
        finally:
            MODULE._text_paragraphs = original

        self.assertEqual(result["summary"]["unchanged"], 1)
        self.assertGreaterEqual(result["summary"]["modified"], 1)
        self.assertEqual(len(result["differences"]), 3)

    # @lat: [[feature-workspace#Contract comparison#Structured Word tables]]
    def test_docx_tables_preserve_column_and_row_merges(self):
        xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body><w:tbl><w:tblGrid><w:gridCol w:w="1000"/><w:gridCol w:w="2000"/></w:tblGrid>
    <w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>标题</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>纵向合并</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>第一行</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc><w:tc><w:p><w:r><w:t>第二行</w:t></w:r></w:p></w:tc></w:tr>
  </w:tbl></w:body>
</w:document>"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/document.xml", xml)
            document = MODULE._docx_document(path, "old")

        table = document["blocks"][0]
        self.assertEqual(table["columnWidths"], [1000, 2000])
        self.assertEqual(table["rows"][0]["cells"][0]["colSpan"], 2)
        self.assertEqual(table["rows"][1]["cells"][0]["rowSpan"], 2)
        self.assertEqual(len(table["rows"][2]["cells"]), 1)

    # @lat: [[feature-workspace#Optional AI interpretation#Word export]]
    def test_contract_analysis_export_converts_markdown_to_docx_structure(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "分析报告.docx"
            result = MODULE._export_contract_analysis(
                {
                    "outputPath": str(output_path),
                    "analysis": (
                        "# 重要变化\n\n"
                        "1. **合同金额**由 15 万元调整为 113.8 万元。\n\n"
                        "| 项目 | 建议 |\n| --- | --- |\n| 金额 | 核对含税口径 |"
                    ),
                    "oldFileName": "旧合同.docx",
                    "newFileName": "新合同.pdf",
                    "modelName": "Claude Opus 5",
                    "perspective": "neutral",
                }
            )
            document = Document(output_path)
            paragraph_text = [paragraph.text for paragraph in document.paragraphs]

            self.assertEqual(result["path"], str(output_path))
            self.assertIn("合同对比 AI 分析报告", paragraph_text)
            self.assertIn("AI 分析结果", paragraph_text)
            self.assertFalse(any("**" in text or text.startswith("#") for text in paragraph_text))
            self.assertTrue(
                any(
                    run.bold and "合同金额" in run.text
                    for paragraph in document.paragraphs
                    for run in paragraph.runs
                )
            )
            self.assertEqual(len(document.tables), 2)

    # @lat: [[work-records#Single record Word export]]
    def test_work_record_export_converts_result_markdown_to_docx(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "工作记录.docx"
            result = MODULE._export_work_record(
                {
                    "outputPath": str(output_path),
                    "title": "合同差异分析",
                    "profileName": "向永驿",
                    "createdAt": 1_789_444_800_000,
                    "status": "completed",
                    "prompt": "请分析两份合同的区别",
                    "steps": [
                        {"label": "读取两份合同", "status": "completed"},
                        {"label": "生成差异摘要", "status": "completed"},
                    ],
                    "resultSummary": (
                        "## 重要变化\n\n"
                        "1. **合同金额**发生调整。\n\n"
                        "| 项目 | 结果 |\n| --- | --- |\n| 金额 | 需要复核 |"
                    ),
                }
            )
            document = Document(output_path)
            paragraph_text = [paragraph.text for paragraph in document.paragraphs]

            self.assertEqual(result["path"], str(output_path))
            self.assertIn("合同差异分析", paragraph_text)
            self.assertIn("我的要求", paragraph_text)
            self.assertIn("执行过程", paragraph_text)
            self.assertIn("结果摘要", paragraph_text)
            self.assertFalse(any("**" in text or text.startswith("#") for text in paragraph_text))
            self.assertEqual(len(document.tables), 2)


if __name__ == "__main__":
    unittest.main()
